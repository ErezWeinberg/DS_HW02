from docx import Document
from docx.shared import Pt
import sys

def main():
    doc = Document()
    doc.add_heading('מבני נתונים 234218 - תרגיל רטוב 2 (חלק יבש)', 0)

    doc.add_heading('1. מבני נתונים', level=1)
    doc.add_paragraph('הפתרון משתמש בשלושה מבני נתונים מרכזיים:')
    doc.add_paragraph('א. טבלת גיבוב (Hash Table):', style='List Bullet')
    doc.add_paragraph('ממפה מזהה מתמודד (contestantId) לצומת במבנה Union-Find. הטבלה משתמשת במיעון פתוח (Open Addressing) עם ניקור ליניארי (Linear Probing). כיוון שמתמודדים לא מוסרים מהמערכת, אין צורך בניהול מחיקות (Tombstones), ופעולות הכנסה וחיפוש לוקחות תוחלת זמן של O(1).')
    
    doc.add_paragraph('ב. עצי AVL (Active Teams):', style='List Bullet')
    doc.add_paragraph('אנו מנהלים שני עצי AVL לקבוצות פעילות. העץ הראשון (teamsById) מסודר לפי מזהה קבוצה (teamId) ומאפשר חיפוש, הכנסה ומחיקה בזמן O(log k). העץ השני (teamsByMotivation) מסודר לפי המוטיבציה הכוללת של הקבוצה (totalMotivation), וכשובר שוויון לפי teamId. עץ זה מתחזק גודל תת-עץ (size) בכל צומת על מנת לאפשר שאילתות סטטיסטיות (בחירת האיבר ה-i) בזמן O(log k).')

    doc.add_paragraph('ג. מבנה קבוצות זרות (Union-Find) עצל:', style='List Bullet')
    doc.add_paragraph('מבנה זה מייצג את הקבוצות ומתחזק שני ערכים בצורה עצלה: מספר משימות (missions_offset) והמיומנות הכרונולוגית (skill_multiplier). לכל צומת יש מצביע לאב. שורש העץ שומר דגל is_active שמסמן אם הקבוצה לא הודחה. ערכים בצומת השורש הם אבסולוטיים (מייצגים את הערך של הקבוצה כולה), בעוד שערכים של ילדים הם יחסיים לשורש. פעולת ה-Union משלבת עצים לפי גודל, ופעולת ה-Find משתמשת בכיווץ מסלולים (Path Compression) תוך עדכון הערכים היחסיים למניעת ספירה כפולה. זמן ממוצע לפעולה הוא O(log* n).')

    doc.add_heading('2. אלגוריתמים וסיבוכיות זמן', level=1)
    
    doc.add_heading('Racenion()', level=2)
    doc.add_paragraph('אתחול מבני הנתונים הריקים: עצי ה-AVL וטבלת הגיבוב. סיבוכיות: O(1).')

    doc.add_heading('~Racenion()', level=2)
    doc.add_paragraph('מעבר על כל טבלת הגיבוב למחיקת צומתי ה-Union-Find, ומעבר על העץ teamsById למחיקת אובייקטי ה-Team. מחיקת מבני העצים. סיבוכיות: O(n + k).')

    doc.add_heading('add_team', level=2)
    doc.add_paragraph('יצירת אובייקט Team והכנסתו לשני עצי ה-AVL. כיוון שההכנסה לעץ AVL היא ב-O(log k), הזמן הכולל הוא O(log k) במקרה הגרוע.')

    doc.add_heading('remove_team', level=2)
    doc.add_paragraph('חיפוש הקבוצה ב-teamsById, מחיקתה משני עצי ה-AVL, ושינוי דגל is_active של השורש שלה ב-Union-Find ל-false (O(1)). מחיקת אובייקט הקבוצה. הכל מבוצע ב-O(log k) זמן מקסימלי בגלל פעולות ה-AVL.')

    doc.add_heading('add_contestant', level=2)
    doc.add_paragraph('חיפוש מתמודד בטבלת הגיבוב ב-O(1) ממוצע. מציאת הקבוצה בעץ ה-AVL ב-O(log k). יצירת צומת חדש והוספתו לטבלת הגיבוב. הוצאת הקבוצה מעץ המוטיבציה, עדכון ה-totalMotivation, והכנסתה חזרה (O(log k)). לאחר מכן מיזוג הצומת כעץ בודד לתוך עץ הקבוצה ב-Union-Find (O(1)). סך הכל: משוערך בממוצע O(log k).')

    doc.add_heading('duel', level=2)
    doc.add_paragraph('חיפוש הקבוצות בעץ ה-AVL ב-O(log k). חישוב הניקוד באמצעות שדות הקבוצה (המתעדכנים תמיד) מתבצע ב-O(1). עדכון ה-teamExp וה-missions_offset בשורשי ה-Union-Find לוקח O(1). סך הכל O(log k) במקרה הגרוע.')

    doc.add_heading('get_contestant_missions_number', level=2)
    doc.add_paragraph('חיפוש בטבלת הגיבוב ב-O(1) ממוצע. מציאת השורש ב-Union-Find בעזרת כיווץ מסלולים ב-O(log* n) משוערך. חישוב הערך המדויק על ידי חיבור ההפרש מהצומת לשורש. זמן כולל משוערך O(log* n) בממוצע.')

    doc.add_heading('get_team_experience', level=2)
    doc.add_paragraph('חיפוש מזהה הקבוצה בעץ ה-AVL והחזרת הערך. סיבוכיות: O(log k) במקרה הגרוע.')

    doc.add_heading('get_ith_collective_motivation_team', level=2)
    doc.add_paragraph('שימוש בשדה ה-size בעץ עצי ה-AVL על מנת למצוא את האיבר ה-i בזמן O(log k) במקרה הגרוע.')

    doc.add_heading('get_partial_team_skill', level=2)
    doc.add_paragraph('שליפת הצומת מטבלת הגיבוב ב-O(1) ממוצע. מציאת השורש ב-Union-Find ב-O(log* n). אם השורש לא פעיל (is_active == false), מוחזרת שגיאה. אחרת, הערך מחושב על ידי הכפלת סך המיומנות של השורש בערך היחסי שבצומת המתמודד ב-O(1). הזמן הכולל הוא O(log* n) משוערך בממוצע.')

    doc.add_heading('recruit', level=2)
    doc.add_paragraph('מציאת הקבוצות והוצאתן משני עצי ה-AVL. חישוב המוטיבציה החדשה של הקבוצה המגייסת והכנסתה לעץ ב-O(log k). מיזוג עצים (union_trees) מתבצע על ידי עדכון אב וחישוב המכפלה החדשה לפי ה-inv() מתוך משפט פרמה בזמן O(1) בלבד. זמן זה מצטרף לשליפות כך שהזמן הכולל הוא O(log k + log* n) משוערך (תלוי במספר הקבוצות והמתמודדים).')

    doc.add_heading('3. סיבוכיות מקום', level=1)
    doc.add_paragraph('עבור n מתמודדים שנוספו למערכת, קיימים לכל היותר n צמתים בטבלת הגיבוב ו-n צמתים ב-Union-Find. המקום הכולל הוא O(n).')
    doc.add_paragraph('עבור k קבוצות פעילות, קיימים k צמתים בכל אחד משני עצי ה-AVL, המקום הכולל הוא O(k).')
    doc.add_paragraph('לכן סיבוכיות המקום הדרושה למבנה הנתונים היא במקרה הגרוע O(n + k) כנדרש.')

    doc.save('DS_HW02_dry.docx')

if __name__ == '__main__':
    main()
